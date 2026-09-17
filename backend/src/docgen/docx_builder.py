"""Word 보고서 생성 — 상세 분석본."""
from __future__ import annotations

import io
from datetime import datetime
from pathlib import Path
from typing import Dict, List

import pandas as pd
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from typing import Any, Optional

from ..aggregation import by_category, missing_departments, total
from ..sensing import anomalies as get_anomalies, summary_counts

KOREAN_FONT = "맑은 고딕"


def _set_run(run, *, size=11, bold=False, color: RGBColor | None = None):
    run.font.name = KOREAN_FONT
    # 한글 폰트 동방언어 적용
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        from docx.oxml.ns import OxmlElement
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    rFonts.set(qn("w:eastAsia"), KOREAN_FONT)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def _add_para(doc: Document, text: str, *, size=11, bold=False, align=None, color=None):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    _set_run(r, size=size, bold=bold, color=color)
    return p


def _add_table(doc: Document, df: pd.DataFrame, headers: List[str], col_widths_cm: List[float]):
    tbl = doc.add_table(rows=len(df) + 1, cols=len(headers))
    tbl.style = "Light Grid Accent 1"
    for j, (h, w) in enumerate(zip(headers, col_widths_cm)):
        c = tbl.rows[0].cells[j]
        c.text = h
        c.width = Cm(w)
        for r in c.paragraphs[0].runs:
            _set_run(r, size=10, bold=True)
    for i, row in df.iterrows():
        for j, h in enumerate(headers):
            v = row[h] if h in row.index else ""
            if isinstance(v, float):
                v = f"{v:.1f}"
            c = tbl.rows[i + 1].cells[j]
            c.text = str(v)
            c.width = Cm(col_widths_cm[j])
            for r in c.paragraphs[0].runs:
                _set_run(r, size=10)


def build_report(*, df: pd.DataFrame, period: str, narrative_summary: str,
                 anomaly_narratives: Dict[int, str],
                 category_insights: Optional[List[Dict[str, str]]] = None,
                 recommendations: Optional[List[str]] = None,
                 clustering: Optional[Dict[str, Any]] = None,
                 strategy: Optional[Dict[str, Any]] = None) -> bytes:
    doc = Document()

    # 표지 제목
    _add_para(doc, f"{period} 경영성과 상세 보고서", size=22, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
    _add_para(doc, "키친 사업부 · Early Sensing Agent", size=12,
              align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x64, 0x74, 0x8B))
    _add_para(doc, f"자동 생성: {datetime.now().strftime('%Y-%m-%d %H:%M')}",
              size=10, align=WD_ALIGN_PARAGRAPH.CENTER, color=RGBColor(0x94, 0xA3, 0xB8))
    doc.add_paragraph()

    # 1. Executive Summary
    _add_para(doc, "1. Executive Summary", size=16, bold=True)
    _add_para(doc, narrative_summary, size=11)

    totals = total(df)
    counts = summary_counts(df)
    _add_para(doc,
              f"• 전체 달성률: {totals['달성률']:.1f}% (목표 {totals['목표']:.0f}억원 / 실적 {totals['실적']:.0f}억원)",
              size=11)
    _add_para(doc,
              f"• 항목 분류: 정상 {counts['정상']} · 경고 {counts['경고']} · 위험 {counts['위험']} (총 {counts['총']})",
              size=11)
    doc.add_paragraph()

    # 2. 영역별 실적 집계
    _add_para(doc, "2. 영역별 실적 집계", size=16, bold=True)
    agg_df = by_category(df).copy()
    agg_df["목표(억원)"] = agg_df["목표(억원)"].astype(float).round(1)
    agg_df["실적(억원)"] = agg_df["실적(억원)"].astype(float).round(1)
    agg_df["달성률(%)"] = agg_df["달성률(%)"].astype(float).round(1)
    _add_table(doc, agg_df, ["카테고리", "목표(억원)", "실적(억원)", "달성률(%)"], [3.5, 3, 3, 3])
    doc.add_paragraph()

    # 2-1. 카테고리별 AI 분석
    if category_insights:
        _add_para(doc, "2-1. 카테고리별 AI 분석", size=14, bold=True)
        for ins in category_insights:
            _add_para(doc, f"• {ins.get('카테고리', '')}: {ins.get('분석', '')}", size=11)
        doc.add_paragraph()

    # 3. 이상 징후
    _add_para(doc, "3. 이상 징후 (Early Sensing)", size=16, bold=True)
    anomalies_df = get_anomalies(df).reset_index(drop=True)
    if anomalies_df.empty:
        _add_para(doc, "모든 항목이 정상 범위 내에 있습니다.", size=11,
                  color=RGBColor(0x10, 0xB9, 0x81))
    else:
        for i, row in anomalies_df.iterrows():
            color = RGBColor(0xEF, 0x44, 0x44) if row["상태"] == "위험" else RGBColor(0xF5, 0x9E, 0x0B)
            _add_para(doc,
                      f"[{row['상태']}] {row['카테고리']} · {row['부서']} · {row['세부항목']} "
                      f"(달성률 {row['달성률(%)']:.1f}%)",
                      size=11, bold=True, color=color)
            narrative = anomaly_narratives.get(i, "")
            if narrative:
                _add_para(doc, f"  · AI 해석: {narrative}", size=10,
                          color=RGBColor(0x64, 0x74, 0x8B))
    doc.add_paragraph()

    # 3-1. 이상 항목 근본원인 클러스터 (모니)
    if clustering and clustering.get("clusters"):
        _add_para(doc, "3-1. 이상 항목 근본원인 클러스터 (모니)", size=14, bold=True)
        _add_para(doc, "[모니 사고 과정]", size=10, bold=True, color=RGBColor(0x64, 0x74, 0x8B))
        for obs in clustering.get("observations", []):
            _add_para(doc, f"  · 관찰: {obs}", size=10, color=RGBColor(0x64, 0x74, 0x8B))
        for i, step in enumerate(clustering.get("reasoning", []), 1):
            _add_para(doc, f"  · 추론 {i}: {step}", size=10, color=RGBColor(0x64, 0x74, 0x8B))
        _add_para(doc, "[클러스터 결과]", size=10, bold=True, color=RGBColor(0x64, 0x74, 0x8B))
        for c in clustering.get("clusters", []):
            item_idxs = c.get("items", [])
            item_labels = []
            for idx in item_idxs:
                if 0 <= idx < len(anomalies_df):
                    r = anomalies_df.iloc[idx]
                    item_labels.append(f"{r['카테고리']}/{r['부서']}")
            _add_para(doc, f"▸ {c.get('name', '')}", size=11, bold=True,
                      color=RGBColor(0xA5, 0x00, 0x34))
            _add_para(doc, f"  · 포함 항목: {', '.join(item_labels) or '(없음)'}", size=10)
            _add_para(doc, f"  · 이유: {c.get('reason', '')}", size=10,
                      color=RGBColor(0x64, 0x74, 0x8B))
        doc.add_paragraph()

    # 4. 미회신 부서
    _add_para(doc, "4. 미회신 부서", size=16, bold=True)
    missing = missing_departments(df)
    if not missing:
        _add_para(doc, "모든 부서가 회신을 완료했습니다.", size=11,
                  color=RGBColor(0x10, 0xB9, 0x81))
    else:
        for cat, depts in missing.items():
            _add_para(doc, f"• {cat}: {', '.join(depts)}", size=11,
                      color=RGBColor(0xEF, 0x44, 0x44))
        _add_para(doc, "→ Follow-up 메일이 자동 발송됩니다.",
                  size=10, color=RGBColor(0x64, 0x74, 0x8B))

    # 4-1. 모니 Follow-up 전략
    if strategy and strategy.get("strategies"):
        doc.add_paragraph()
        _add_para(doc, "4-1. Follow-up 전략 (모니 결정)", size=14, bold=True)
        _add_para(doc, "[모니 사고 과정]", size=10, bold=True, color=RGBColor(0x64, 0x74, 0x8B))
        for obs in strategy.get("observations", []):
            _add_para(doc, f"  · 관찰: {obs}", size=10, color=RGBColor(0x64, 0x74, 0x8B))
        for i, step in enumerate(strategy.get("reasoning", []), 1):
            _add_para(doc, f"  · 추론 {i}: {step}", size=10, color=RGBColor(0x64, 0x74, 0x8B))
        for s in strategy.get("strategies", []):
            pri = s.get("priority", "medium")
            color = (RGBColor(0xEF, 0x44, 0x44) if pri == "high"
                     else RGBColor(0xF5, 0x9E, 0x0B) if pri == "medium"
                     else RGBColor(0x64, 0x74, 0x8B))
            _add_para(doc,
                      f"▸ {s.get('department')} ({s.get('category')}) — "
                      f"전략: {s.get('strategy')} / 우선순위: {pri}",
                      size=11, bold=True, color=color)
            _add_para(doc, f"  · 이유: {s.get('reason', '')}", size=10,
                      color=RGBColor(0x64, 0x74, 0x8B))

    # 5. 권장 액션 (AI)
    if recommendations:
        doc.add_paragraph()
        _add_para(doc, "5. AI 권장 액션 — 차주 우선순위", size=16, bold=True)
        for i, rec in enumerate(recommendations, 1):
            _add_para(doc, f"  {i}. {rec}", size=11)

    out = io.BytesIO()
    doc.save(out)
    out.seek(0)
    return out.read()


def save_report(path: Path, **kwargs) -> Path:
    data = build_report(**kwargs)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(data)
    return path
