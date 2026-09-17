"""LG 실제 양식(ver0.1 창원대공유 'Master추가') 파서.

ver0.2(`당월>>` 마커형)에서 양식이 통째로 바뀜 → 마커 폐기. 새 양식 특징:
- 시트 = Master + Area_A~L (12 영역). Master 는 Area 들의 엑셀 수식 집계 → 파싱 X.
- 각 Area 는 헤더 위치/계층 깊이/주차여부가 제각각 → **고정 행/열 가정 금지, 동적 감지**.
- 한 셀이 **raw 입력**인지 **수식(소계·달성율·Total·롤업)**인지 구분 가능
  (data_only=False 면 수식은 '=' 텍스트). → **raw 말단 셀만 적재**, 수식은 재계산 대상.

핵심 설계 (사용자 합의, 2026-06-19):
- DB 엔 raw 만 = source of truth. 소계/달성율/Master 는 읽을 때 재계산(aggregate_lg).
- 출력 스키마는 구 파서와 **동일**(store/aggregate 무변경): 항목·부서경로·데이터종류·
  연도·종류·월·구분(월계|주차)·주차·값·원본셀.
- 구조 패턴 4종을 한 동적 파서로 흡수:
  ① 주차+월 / 행 dtype (A·B·C·D·G)   ② 연도비교 주차 (E·L)
  ③ 월만 / 행 dtype (F)              ④ 월만 / 연도비교 행 dtype (H·I·J·K)
  → 연도가 행 라벨(`25년 실적)이면 _split_dtype, 상단 컬럼블록(Area_L 25/26)이면 그쪽서.

mock 파서(parsing.py)는 건드리지 않음. 이 모듈은 LG 실양식 전용.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Dict, List, Optional, Tuple, Union

import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from . import registry

FileLike = Union[str, Path]

OUT_COLUMNS = ["항목", "부서경로", "데이터종류", "연도", "종류",
               "월", "구분", "주차", "값", "원본셀"]

# Master/요약 시트는 집계 결과(수식)라 원천 아님 → skip
SKIP_SHEETS = {"Master", "마스터", "Summary", "요약"}

_MONTH_RE = re.compile(r"^\s*(\d{1,2})\s*월\s*$")
_WEEK_RE = re.compile(r"^\s*W\s*(\d{1,2})\s*$", re.IGNORECASE)
_DTYPE_RE = re.compile(r"(목표|실적|계획)")          # 달성율/달성률은 파생 → 안 잡음
# 연도블록 헤더 — 두 표기 모두: '25년 실적'(년) / '`25 실적'(백틱, 년 없음, ver0.1 매출 시트)
_YEARBLOCK_RE = re.compile(r"(?:`\s*\d{2}|\d{2}\s*년)")

MONTHS = [f"{i}월" for i in range(1, 13)]


def _split_dtype(s: str) -> Tuple[Optional[int], str]:
    """데이터종류 원문 → (연도, 종류). 연도 없으면 None.

    '`25년 실적' → (2025, '실적'), '`26년 목표' → (2026, '목표'),
    '목표' → (None, '목표'), '실적' → (None, '실적').
    연도를 코드에 박지 않게 2자리 연도를 동적 추출(2027 등 자동 대응).
    """
    m = re.search(r"(\d{2})", s)
    year = 2000 + int(m.group(1)) if m else None
    for k in ("목표", "실적", "계획"):
        if k in s:
            return year, k
    return year, s.strip()


def _month_label(v) -> Optional[str]:
    if isinstance(v, str):
        m = _MONTH_RE.match(v)
        if m:
            return f"{int(m.group(1))}월"
    return None


def _week_label(v) -> Optional[str]:
    if v is None:
        return None
    m = _WEEK_RE.match(str(v).strip())
    return f"W{int(m.group(1))}" if m else None


def _is_dtype(v) -> bool:
    return isinstance(v, str) and bool(_DTYPE_RE.search(v))


def _find_month_row(wsv: Worksheet, scan_rows: int = 8) -> int:
    """월 토큰(N월)이 가장 많은 상단 행 = 월 헤더 행 (1-based)."""
    best_row, best_cnt = 1, -1
    for r in range(1, min(scan_rows, wsv.max_row) + 1):
        cnt = sum(1 for c in wsv[r] if _month_label(c.value))
        if cnt > best_cnt:
            best_row, best_cnt = r, cnt
    return best_row


def _build_month_map(wsv: Worksheet, month_row: int, max_col: int) -> Dict[int, Optional[str]]:
    """컬럼 → 월 라벨. 병합셀이라 월 시작 컬럼에만 값 → forward-fill."""
    out: Dict[int, Optional[str]] = {}
    cur: Optional[str] = None
    for c in range(1, max_col + 1):
        v = wsv.cell(month_row, c).value
        lab = _month_label(v)
        if lab:
            cur = lab
        elif isinstance(v, str) and v.strip() == "Total":
            # 연간 Total 컬럼(월헤더 자체가 'Total') = 블록 끝 → 이후 직전 월 전파 끊음.
            # (Area_L 처럼 25블록 Total 뒤 26블록 1월이 새로 시작하게)
            cur = None
            out[c] = None
            continue
        out[c] = cur
    return out


def _build_gran_map(wsv: Worksheet, month_row: int, max_col: int) -> Dict[int, Tuple[str, Optional[str]]]:
    """컬럼 → (구분, 주차). 월헤더행 + 바로 아래행에서 실적/W#/Total 판정.

    - 'W#'  → ('주차', 'W#')
    - '실적' / 'Total'(월 밑 = 그 달 월값) / 그 외 단일 컬럼 → ('월계', None)
    연간 Total 컬럼은 _build_month_map 에서 이미 제외(월=None) → 여기 안 옴.
    """
    out: Dict[int, Tuple[str, Optional[str]]] = {}
    for c in range(1, max_col + 1):
        kind, wl = "월계", None
        for r in (month_row, month_row + 1):
            wlab = _week_label(wsv.cell(r, c).value)
            if wlab:
                kind, wl = "주차", wlab
                break
        out[c] = (kind, wl)
    return out


def _find_year_blocks(wsv: Worksheet, month_row: int, max_col: int) -> Dict[int, Tuple[Optional[int], Optional[str]]]:
    """상단(월헤더 위쪽) 연도블록 헤더 → 컬럼별 (연도, 종류). 없으면 모두 (None, None).

    Area_L 처럼 '25년 실적' / '26년 목표' 가 컬럼 블록(상단 병합)일 때 사용.
    forward-fill 로 블록 범위 전체에 전파. 월헤더 다음 컬럼부터 다음 블록 전까지.
    """
    out: Dict[int, Tuple[Optional[int], Optional[str]]] = {c: (None, None) for c in range(1, max_col + 1)}
    if month_row <= 1:
        return out
    found = False
    cur: Tuple[Optional[int], Optional[str]] = (None, None)
    for c in range(1, max_col + 1):
        for r in range(1, month_row):
            v = wsv.cell(r, c).value
            if isinstance(v, str) and _YEARBLOCK_RE.search(v) and _is_dtype(v):
                yr, kind = _split_dtype(v)
                cur = (yr, kind)
                found = True
                break
        out[c] = cur
    return out if found else {c: (None, None) for c in range(1, max_col + 1)}


def parse_sheet(wsv: Worksheet, wsf: Worksheet) -> pd.DataFrame:
    """Area 시트 한 장(값 wsv + 수식 wsf) → tidy DataFrame (raw 말단 셀만).

    동적 헤더 감지 → 좌측 라벨 ffill + dtype 열 자동 → 각 데이터 셀이
    수식이면 skip(소계/달성율/Total/롤업), raw 숫자면 1레코드.
    """
    max_col = wsv.max_column
    max_row = wsv.max_row
    month_row = _find_month_row(wsv)
    month_map = _build_month_map(wsv, month_row, max_col)
    gran_map = _build_gran_map(wsv, month_row, max_col)
    year_blocks = _find_year_blocks(wsv, month_row, max_col)
    has_year_block = any(yb[0] for yb in year_blocks.values())

    # 데이터 컬럼 = 월 라벨이 매핑된 컬럼. 그 왼쪽 = 라벨 영역.
    data_cols = [c for c in range(1, max_col + 1) if month_map.get(c)]
    if not data_cols:
        return pd.DataFrame(columns=OUT_COLUMNS)
    first_data_col = min(data_cols)
    label_area = list(range(1, first_data_col))
    data_start = month_row + 1

    # dtype 열: 라벨 영역 중 목표/실적/계획 매칭이 가장 많은 열 (행 dtype 시트)
    hits = {c: 0 for c in label_area}
    for r in range(data_start, max_row + 1):
        for c in label_area:
            if _is_dtype(wsv.cell(r, c).value):
                hits[c] += 1
    dtype_col = max(hits, key=lambda c: hits[c]) if hits and max(hits.values()) > 0 else None
    label_cols = [c for c in label_area if c != dtype_col]
    label_state: Dict[int, Optional[str]] = {c: None for c in label_cols}
    # dtype(목표/실적 + 연도) ffill 상태 — 투자비·매출2 처럼 블록당 한 번만 적힌 양식 대응.
    # master_lg 평가기(dt_year/dt_kind)와 동일 처리 → raw CSV 가 수식 평가본과 정합
    # (블록 연속행의 실적 셀이 빈 dtype 으로 스킵돼 누락되던 것 방지). 리셋 없음(다음 헤더가 갱신).
    dt_year: Optional[int] = None
    dt_kind: str = ""
    dt_raw: str = ""

    records: List[dict] = []
    title = wsv.title
    for r in range(data_start, max_row + 1):
        # 좌측 라벨 forward-fill
        for c in label_cols:
            v = wsv.cell(r, c).value
            if v is not None and str(v).strip():
                label_state[c] = str(v).strip()
        path = " / ".join(x for x in label_state.values() if x)

        # 행 dtype (목표/실적/`25년실적…) — 빈 칸은 직전 dtype 헤더 상속(ffill). 없으면 연도블록에서 채움.
        if dtype_col is not None:
            dv = wsv.cell(r, dtype_col).value
            if dv is not None and str(dv).strip():
                dt_raw = str(dv).strip()
                dt_year, dt_kind = _split_dtype(dt_raw)
            dtype_raw, row_year, row_kind = dt_raw, dt_year, dt_kind
        else:
            dtype_raw, row_year, row_kind = "", None, ""

        for c in data_cols:
            fval = wsf.cell(r, c).value
            if isinstance(fval, str) and fval.startswith("="):
                continue  # 수식 셀 = 소계/달성율/Total/롤업 → 재계산 대상, 적재 X
            val = wsv.cell(r, c).value
            if val is None or isinstance(val, str):
                continue  # 빈칸 / #DIV/0! 등
            try:
                num = float(val)
            except (TypeError, ValueError):
                continue

            gran, week = gran_map[c]

            # 연도/종류 결정: 행 dtype 우선, 없으면 컬럼 연도블록
            yb_year, yb_kind = year_blocks.get(c, (None, None))
            kind = row_kind or yb_kind
            if not kind:
                continue  # 목표/실적 구분 없는 셀 = 헤더 잔재/보조값
            year = row_year if row_year is not None else yb_year

            records.append({
                "항목": title,
                "부서경로": path,
                "데이터종류": dtype_raw or (f"`{str(yb_year)[2:]}년 {yb_kind}" if yb_year else kind),
                "연도": year,
                "종류": kind,
                "월": int(month_map[c][:-1]),
                "구분": "주차" if gran == "주차" else "월계",
                "주차": week,
                "값": num,
                "원본셀": f"{title}!{get_column_letter(c)}{r}",
            })
    return pd.DataFrame(records, columns=OUT_COLUMNS)


# 마스터 양식 필수 조건 (빡센 검증) — 부서 제출 = 이 양식 아니면 거부
# 시트 이름은 레지스트리(단일 진실원천)에서. LG 가 이름 바꾸면 registry.py 만 수정.
EXPECTED_AREA_SHEETS = list(registry.SHEET_NAMES)  # 12개 항목 시트
MIN_AREA_SHEETS = 8      # 최소 이만큼 항목 시트 있어야 마스터 양식 인정
MIN_PARSED_ROWS = 100    # 최소 이만큼 데이터 행 나와야 (빈/엉뚱 파일 거부)


def validate_master_format(path_or_bytes: FileLike) -> dict:
    """업로드 파일이 **마스터 양식**인지 빡세게 검증.

    반환: {ok: bool, reason: str, area_sheets: [...], rows: int, missing: [...]}
    조건: (1) Area_* 시트가 MIN_AREA_SHEETS 이상 (2) Master 시트 존재
          (3) 파싱 시 데이터 행 MIN_PARSED_ROWS 이상 (4) 목표·실적 종류 둘 다 존재.
    실패 시 ok=False + 사유. (운영자 원천 = 검증 안 함, 이 함수 안 거침)
    """
    try:
        if isinstance(path_or_bytes, (str, Path)):
            wb = load_workbook(path_or_bytes, read_only=True, data_only=True)
        else:
            import io
            raw = path_or_bytes.read() if hasattr(path_or_bytes, "read") else path_or_bytes
            wb = load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
    except Exception as e:
        return {"ok": False, "reason": f"엑셀 파일을 열 수 없습니다: {e}",
                "area_sheets": [], "rows": 0, "missing": EXPECTED_AREA_SHEETS}

    sheets = set(wb.sheetnames)
    area_sheets = [s for s in EXPECTED_AREA_SHEETS if s in sheets]
    missing = [s for s in EXPECTED_AREA_SHEETS if s not in sheets]
    has_master = bool(sheets & SKIP_SHEETS)

    if len(area_sheets) < MIN_AREA_SHEETS:
        return {"ok": False,
                "reason": f"마스터 양식이 아닙니다 — Area 시트가 {len(area_sheets)}개뿐 "
                          f"(최소 {MIN_AREA_SHEETS}개 필요). 부서 기입용 마스터 엑셀을 올려주세요.",
                "area_sheets": area_sheets, "rows": 0, "missing": missing}
    if not has_master:
        return {"ok": False,
                "reason": "마스터 양식이 아닙니다 — Master(집계) 시트가 없습니다.",
                "area_sheets": area_sheets, "rows": 0, "missing": missing}

    # 실제 파싱해서 데이터가 나오는지 + 목표/실적 둘 다 있는지
    try:
        df, _ = parse_workbook(path_or_bytes)
    except Exception as e:
        return {"ok": False, "reason": f"양식 구조를 해석할 수 없습니다: {e}",
                "area_sheets": area_sheets, "rows": 0, "missing": missing}

    kinds = set(df["종류"].dropna().unique()) if len(df) else set()
    if len(df) < MIN_PARSED_ROWS:
        return {"ok": False,
                "reason": f"데이터가 너무 적습니다 ({len(df)}행, 최소 {MIN_PARSED_ROWS}행). "
                          "빈 양식이거나 다른 파일일 수 있습니다.",
                "area_sheets": area_sheets, "rows": int(len(df)), "missing": missing}
    if not ({"목표", "계획"} & kinds) or "실적" not in kinds:
        return {"ok": False,
                "reason": "목표/실적 데이터가 모두 있어야 합니다 (한쪽만 감지됨).",
                "area_sheets": area_sheets, "rows": int(len(df)), "missing": missing}

    return {"ok": True, "reason": "마스터 양식 확인",
            "area_sheets": area_sheets, "rows": int(len(df)), "missing": missing}


def parse_workbook(path_or_bytes: FileLike) -> Tuple[pd.DataFrame, Dict[str, dict]]:
    """LG 통합 워크북 → (tidy DataFrame, 시트별 진단 dict).

    Master/요약 시트는 skip(집계 수식). Area_* 만 raw 추출.
    값/수식 두 번 로드해 raw vs 수식 구분.
    """
    # 값/수식 두 번 로드. 바이트 스트림은 한 번 읽으면 소진되므로 버퍼링 후 각각 로드.
    # 이미 로드된 (wb_v, wb_f) 쌍을 받으면 재사용 — 같은 파일을 CSV 적재(여기)와 캐시 빌더가
    # 각각 열면 268KB xlsx 를 4번 파싱(openpyxl 로드가 비용 대부분)하므로, save_excel 이 1회
    # 로드해 공유시킨다(읽기 전용이라 안전).
    if isinstance(path_or_bytes, tuple) and len(path_or_bytes) == 2:
        wb_v, wb_f = path_or_bytes
    elif isinstance(path_or_bytes, (str, Path)):
        wb_v = load_workbook(path_or_bytes, data_only=True)
        wb_f = load_workbook(path_or_bytes, data_only=False)
    else:
        import io
        raw = path_or_bytes.read() if hasattr(path_or_bytes, "read") else path_or_bytes
        wb_v = load_workbook(io.BytesIO(raw), data_only=True)
        wb_f = load_workbook(io.BytesIO(raw), data_only=False)

    frames: List[pd.DataFrame] = []
    report: Dict[str, dict] = {}
    for wsv in wb_v.worksheets:
        title = wsv.title
        if title in SKIP_SHEETS:
            report[title] = {"status": "skip", "reason": "집계 시트(수식) — 원천 아님"}
            continue
        slug = registry.slug_for_sheet(title)
        if slug is None:
            report[title] = {"status": "skip", "reason": "레지스트리 미등록 시트 — 항목 아님"}
            continue
        wsf = wb_f[title]
        df = parse_sheet(wsv, wsf)
        if len(df):
            df["항목"] = slug  # 시트명(한글) → 내부 슬러그로 정규화 (단일 진실원천)
        report[title] = {
            "status": "ok" if len(df) else "empty",
            "rows": int(len(df)),
            "데이터종류": sorted(df["데이터종류"].dropna().unique().tolist()) if len(df) else [],
            "월수": int(df["월"].nunique()) if len(df) else 0,
            "주차행": int((df["구분"] == "주차").sum()) if len(df) else 0,
        }
        if len(df):
            frames.append(df)

    if not frames:
        return pd.DataFrame(columns=OUT_COLUMNS), report
    combined = pd.concat(frames, ignore_index=True)

    # 연도 보정: 행/블록 둘 다 연도 없는 셀(A~G 부서제출분)은 파일 기준연도로.
    known = combined["연도"].dropna()
    if len(known):
        base_year = int(known.max())
    else:
        base_year = 2026  # 명시 연도 전무 시 (안전 기본값)
    combined["연도"] = pd.to_numeric(combined["연도"], errors="coerce").fillna(base_year).astype(int)
    combined["연도구분"] = combined["연도"].map(
        lambda y: "당해" if y == base_year else ("전년" if y < base_year else "차기")
    )
    return combined, report


if __name__ == "__main__":  # 간이 점검: python -m src.parsing_lg <xlsx>
    import sys

    sys.stdout.reconfigure(encoding="utf-8")
    src = sys.argv[1] if len(sys.argv) > 1 else \
        "../260619_데이터양식_창원대공유_항목수치랜덤화_Master추가_ver0.1.xlsx"
    df, rep = parse_workbook(src)
    print(f"총 {len(df)} 레코드\n")
    for sheet, info in rep.items():
        print(f"  [{sheet}] {info}")
    print("\n--- 샘플 (Area_A 상위 10행) ---")
    print(df[df["항목"] == "Area_A"].head(10).to_string(index=False))
